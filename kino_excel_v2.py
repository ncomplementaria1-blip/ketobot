"""
Genera Excels de Kino con el formato del usuario:
  Sorteo | Fecha | (1..5) | NEGRO | (6..10) | NEGRO | (11..15) | NEGRO | (16..20) | NEGRO | (21..25)

Cada número va en su columna con su color; las celdas sin número quedan en blanco.
Las columnas separadoras son completamente negras.

Uso:
    python kino_excel_v2.py
    -> regenera analisis_kino/kino_3235.xlsx con los datos actuales y las
       sugerencias para el próximo sorteo (3235).
"""
import json
import os
import random
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "analisis_kino", "kino_data.json")
OUT = os.path.join(BASE, "analisis_kino", "kino_3235.xlsx")

PALETTE = {
    1: ("B8D8E8", "000000"), 2: ("E63946", "000000"), 3: ("A8E6A3", "000000"),
    4: ("F2C879", "000000"), 5: ("7B3FA0", "FFFFFF"), 6: ("1ABC9C", "000000"),
    7: ("FF8FCF", "000000"), 8: ("2980B9", "FFFFFF"), 9: ("7F8C8D", "000000"),
    10: ("A98E33", "FFFFFF"), 11: ("F1C40F", "000000"), 12: ("C8A2D8", "000000"),
    13: ("1ABEDB", "000000"), 14: ("5DADE2", "000000"), 15: ("F39C12", "000000"),
    16: ("1A1A1A", "FFFFFF"), 17: ("EC407A", "FFFFFF"), 18: ("C5E1A5", "000000"),
    19: ("8B7500", "FFFFFF"), 20: ("C0392B", "FFFFFF"), 21: ("00BCD4", "000000"),
    22: ("BDBDBD", "000000"), 23: ("4DD0E1", "000000"), 24: ("D7B98E", "FFFFFF"),
    25: ("F8BBD0", "000000"),
}
BLACK = PatternFill(start_color="000000", end_color="000000", fill_type="solid")
CENTER = Alignment(horizontal="center", vertical="center")
SIDE = Side(style="thin", color="888888")
BORDER = Border(left=SIDE, right=SIDE, top=SIDE, bottom=SIDE)

# Mapeo: número 1..25 -> columna en la hoja (después de Sorteo+Fecha)
# Estructura: cols 3..7 (1..5), 8 NEGRO, 9..13 (6..10), 14 NEGRO,
#             15..19 (11..15), 20 NEGRO, 21..25 (16..20), 26 NEGRO, 27..31 (21..25)
def col_for(n):
    return 3 + 6 * ((n - 1) // 5) + ((n - 1) % 5)

SEPARATOR_COLS = [8, 14, 20, 26]
LAST_COL = 31


def style_number(cell, n):
    bg, fg = PALETTE[n]
    cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    cell.font = Font(bold=True, color=fg)
    cell.alignment = CENTER
    cell.border = BORDER


def paint_separator(ws, row):
    for c in SEPARATOR_COLS:
        ws.cell(row=row, column=c).fill = BLACK


def write_header(ws, row=1):
    ws.cell(row=row, column=1, value="Sorteo").font = Font(bold=True)
    ws.cell(row=row, column=2, value="Fecha").font = Font(bold=True)
    for n in range(1, 26):
        style_number(ws.cell(row=row, column=col_for(n), value=n), n)
    paint_separator(ws, row)


def write_draw_row(ws, row, sorteo, fecha, nums):
    ws.cell(row=row, column=1, value=sorteo).alignment = CENTER
    ws.cell(row=row, column=2, value=str(fecha) if fecha else "").alignment = CENTER
    for n in nums:
        style_number(ws.cell(row=row, column=col_for(n), value=n), n)
    paint_separator(ws, row)


def set_widths(ws):
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 12
    for c in range(3, LAST_COL + 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 4
    # separadoras más finas
    for c in SEPARATOR_COLS:
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 1.5


def build(draws):
    freq = Counter()
    for d in draws:
        for n in d["nums"]:
            freq[n] += 1
    total = len(draws)
    gap = {}
    for n in range(1, 26):
        g = 0
        for d in reversed(draws):
            if n in d["nums"]:
                break
            g += 1
        gap[n] = g
    recent = Counter()
    for d in draws[-10:]:
        for n in d["nums"]:
            recent[n] += 1
    return freq, gap, recent, total


def sheet_sugerencias(wb, draws, freq, gap, recent):
    ws = wb.active
    ws.title = "Sugerencias 3235"
    ws["A1"] = "Sugerencias para sorteo 3235 — referencia (NO predice)"
    ws["A1"].font = Font(bold=True, size=13)
    ranked = sorted(range(1, 26), key=lambda n: -freq[n])
    atras = sorted(range(1, 26), key=lambda n: (-gap[n], -freq[n]))
    trend = sorted(range(1, 26), key=lambda n: (-recent[n], -freq[n]))
    random.seed(3235)
    combos = [
        ("Calientes históricos", sorted(ranked[:14])),
        ("Frías históricas", sorted(ranked[-14:])),
        ("Atrasados", sorted(atras[:14])),
        ("Tendencia últimos 10", sorted(trend[:14])),
        ("Mixta atrasados+tendencia", sorted(set(atras[:7] + trend[:7]))),
        ("Azar puro", sorted(random.sample(range(1, 26), 14))),
    ]
    # encabezado con números
    write_header(ws, row=3)
    ws.cell(row=3, column=1, value="Estrategia").font = Font(bold=True)
    for i, (label, nums) in enumerate(combos, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        for n in nums:
            style_number(ws.cell(row=i, column=col_for(n), value=n), n)
        paint_separator(ws, i)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 4
    for c in range(3, LAST_COL + 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 4
    for c in SEPARATOR_COLS:
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = 1.5


def sheet_sorteos(wb, draws):
    ws = wb.create_sheet("Sorteos ordenados")
    write_header(ws, row=1)
    for i, d in enumerate(draws, start=2):
        write_draw_row(ws, i, d["sorteo"], d.get("fecha", ""), d["nums"])
    set_widths(ws)
    ws.freeze_panes = "C2"


def sheet_ultimos(wb, draws):
    ws = wb.create_sheet("Últimos 20")
    ws["A1"] = f"Últimos 20 sorteos ({draws[-20]['sorteo']}–{draws[-1]['sorteo']})"
    ws["A1"].font = Font(bold=True, size=12)
    write_header(ws, row=3)
    for i, d in enumerate(draws[-20:], start=4):
        write_draw_row(ws, i, d["sorteo"], d.get("fecha", ""), d["nums"])
    set_widths(ws)


def sheet_frecuencias(wb, freq, gap, recent, total):
    ws = wb.create_sheet("Frecuencias")
    ws["A1"] = f"Frecuencias — {total} sorteos"
    ws["A1"].font = Font(bold=True, size=12)
    for c, h in enumerate(["N°", "Veces", "%", "Atraso", "Últimos 10"], start=1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = Font(bold=True); cell.alignment = CENTER; cell.border = BORDER
    ranked = sorted(range(1, 26), key=lambda n: -freq[n])
    for i, n in enumerate(ranked, start=4):
        style_number(ws.cell(row=i, column=1, value=n), n)
        ws.cell(row=i, column=2, value=freq[n]).alignment = CENTER
        ws.cell(row=i, column=3, value=round(freq[n] / total * 100, 2)).alignment = CENTER
        ws.cell(row=i, column=4, value=gap[n]).alignment = CENTER
        ws.cell(row=i, column=5, value=recent[n]).alignment = CENTER
    for col, w in zip("ABCDE", (8, 10, 10, 12, 14)):
        ws.column_dimensions[col].width = w


def main():
    with open(DATA, encoding="utf-8") as f:
        draws = json.load(f)
    freq, gap, recent, total = build(draws)
    wb = Workbook()
    sheet_sugerencias(wb, draws, freq, gap, recent)
    sheet_sorteos(wb, draws)
    sheet_ultimos(wb, draws)
    sheet_frecuencias(wb, freq, gap, recent, total)
    wb.save(OUT)
    print(f"OK -> {OUT} ({total} sorteos)")


if __name__ == "__main__":
    main()
